"""Read career-relevant text from an uploaded resume. The file is not stored."""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARS = 10_000
MIN_RESUME_CHARS = 80
MAX_PDF_PAGES = 15

ALLOWED_EXTENSIONS = frozenset({"pdf", "docx"})


class ResumeReadError(Exception):
    """User-facing resume read failure. status_code is an HTTP status."""

    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def file_extension(filename: str | None) -> str:
    if not filename or "." not in filename:
        return ""
    return filename.rsplit(".", 1)[-1].lower()


_GLUED_JUNK_RE = re.compile(r"(?:\.is[A-Z]\w*|(?:\.0){3,})")


def normalize_resume_text(text: str) -> str:
    cleaned = text.replace("\x00", " ")
    cleaned = _GLUED_JUNK_RE.sub(" ", cleaned)
    cleaned = re.sub(r"[^\S\n]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(BytesIO(data))
    except PdfReadError as exc:
        raise ResumeReadError(
            "We couldn't read this PDF. Upload a text-based PDF or a DOCX file."
        ) from exc

    if reader.is_encrypted and reader.decrypt("") == 0:
        raise ResumeReadError(
            "This PDF is password-protected. Upload an unlocked PDF or a DOCX file."
        )

    parts: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    if not zipfile.is_zipfile(BytesIO(data)):
        raise ResumeReadError(
            "This Word file could not be read. Save it as DOCX or PDF and try again."
        )

    try:
        document = Document(BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, ValueError) as exc:
        raise ResumeReadError(
            "This Word file could not be read. Save it as DOCX or PDF and try again."
        ) from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = []
            seen: set[str] = set()
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in seen:
                    seen.add(value)
                    cells.append(value)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_resume_text(data: bytes, filename: str | None) -> str:
    """Validate an in-memory resume and return plain text for AI mapping."""
    if not data:
        raise ResumeReadError("The uploaded file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeReadError("Resume must be 5 MB or smaller.", status_code=413)

    extension = file_extension(filename)
    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeReadError("Upload a PDF or DOCX resume.")

    if extension == "pdf":
        if not data.startswith(b"%PDF"):
            raise ResumeReadError("This file is not a valid PDF.")
        raw = _extract_pdf(data)
    else:
        if not data.startswith(b"PK"):
            raise ResumeReadError("This file is not a valid DOCX resume.")
        raw = _extract_docx(data)

    text = normalize_resume_text(raw)
    if len(text) < MIN_RESUME_CHARS:
        raise ResumeReadError(
            "We couldn't read enough text from this resume. "
            "If it is a scanned image, upload a text-based PDF or DOCX, "
            "or describe your background instead."
        )
    return text[:MAX_RESUME_CHARS]
